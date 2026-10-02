"""Read an eToro workbook without modifying it; send extracted rows to stdout.

Run with the bundled Python runtime and pipe output directly to private app data.
Never redirect real statement output into tracked source or fixtures.
"""
import datetime
import json
import sys
import openpyxl

workbook = openpyxl.load_workbook(sys.argv[1], read_only=True, data_only=True)
result = {}
for name in ['Closed Positions', 'Account Activity', 'Dividends']:
    if name not in workbook.sheetnames:
        raise ValueError('Missing required eToro statement sheet: ' + name)
    iterator = workbook[name].iter_rows(values_only=True)
    headers = list(next(iterator))
    result[name] = [dict(zip(headers, row)) for row in iterator if any(v is not None for v in row)]
workbook.close()
json.dump(result, sys.stdout, default=lambda value: value.isoformat()
          if isinstance(value, (datetime.datetime, datetime.date)) else str(value))
